from __future__ import annotations

from pathlib import Path
from app.parsers.ocr import OcrError, image_bytes_text, image_text, render_pdf_page


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return raw.decode("utf-8", errors="ignore")


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    blocks = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _pdf_content(path: Path) -> tuple[str, bool]:
    import pdfplumber

    pages_text = []
    ocr_used = False
    with pdfplumber.open(str(path)) as pdf:
        if len(pdf.pages) > 30:
            raise OcrError("too_many_pdf_pages", 413)
        for index, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            # Image-bearing pages with little native text require OCR. A short,
            # pure text PDF must not depend on Tesseract being installed.
            meaningful = len("".join(text.split()))
            if meaningful == 0 or (meaningful < 80 and page.images):
                with render_pdf_page(path, index) as image:
                    ocr = image_text(image)
                if len("".join(ocr.split())) > meaningful:
                    text = ocr
                ocr_used = True
            pages_text.append(text.strip())
    return "\n\n".join(pages_text), ocr_used


def _pdf_text(path: Path) -> str:
    return _pdf_content(path)[0]


def _ocr_image_bytes(image_bytes: bytes) -> str:
    return image_bytes_text(image_bytes)


def _ocr_pdf(path: Path) -> str:
    import pdfplumber
    texts = []
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > 30:
            raise OcrError("too_many_pdf_pages", 413)
        for index in range(len(pdf.pages)):
            with render_pdf_page(path, index) as image:
                texts.append(image_text(image))
    return "\n\n".join(texts)


def extract_resume_text(path: str | Path) -> dict:
    """Extract text from pdf/docx/txt/image without touching any LLM.

    Embedded photographs are never passed to any model: only document text or
    the OCR result of text-bearing pages/files is returned for user review.
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    method = "text"
    text = ""
    ocr_used = False

    if suffix == ".pdf":
        text, ocr_used = _pdf_content(file_path)
        method = "ocr" if ocr_used else "text"
    elif suffix == ".docx":
        text = _docx_text(file_path)
    elif suffix == ".txt":
        text = _read_text_file(file_path)
    elif suffix in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
        text = _ocr_image_bytes(file_path.read_bytes())
        method = "ocr"
        ocr_used = True
    else:
        raise ValueError(f"Unsupported resume format: {suffix}")

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text.strip():
        raise ValueError("No readable text found in the resume file.")
    return {
        "text": text.strip(),
        "method": method,
        "ocr_used": ocr_used,
        "file_name": file_path.name,
    }


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}


def ocr_image_file(path: str | Path) -> str:
    """OCR any image used as a job description upload."""
    file_path = Path(path)
    return _ocr_image_bytes(file_path.read_bytes()).strip()
